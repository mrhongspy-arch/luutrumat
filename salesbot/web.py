"""Password-protected admin website, served by the bot process.

Pages: overview (revenue), products + stock import, orders (+ Excel export)
and promotional posts sent through the bot, now or at a scheduled time.
"""

import asyncio
import csv
import hmac
import io
import logging
import re
import secrets
from datetime import date, datetime, timedelta, timezone
from html import escape
from urllib.parse import urlencode

from aiohttp import web
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from telegram.error import TelegramError

from .db import Database

log = logging.getLogger(__name__)

COOKIE = "salesbot_session"
PAGE_SIZE = 50
LOW_STOCK = 3
CHART_DAYS = 30
IMAGE_TYPES = {".jpg", ".jpeg", ".png", ".webp"}
STATUS = {
    "pending": ("Chờ thanh toán", "warn"),
    "paid": ("Đã giao", "good"),
    "cancelled": ("Đã huỷ", "muted"),
}
POST_STATUS = {
    "scheduled": ("Đã hẹn giờ", "warn"),
    "sending": ("Đang gửi", "warn"),
    "sent": ("Đã gửi", "good"),
    "cancelled": ("Đã huỷ", "muted"),
}


# ---------------------------------------------------------------- helpers

def money(amount: int | None) -> str:
    return f"{amount or 0:,}".replace(",", ".") + "đ"


def local_time(iso: str | None, fmt: str = "%d/%m/%Y %H:%M") -> str:
    if not iso:
        return ""
    return datetime.fromisoformat(iso).astimezone().strftime(fmt)


def parse_price(value: str) -> int:
    digits = re.sub(r"\D", "", value or "")
    if not digits:
        raise ValueError("Giá không hợp lệ")
    return int(digits)


def customer(order) -> str:
    return f"@{order['username']}" if order["username"] else str(order["user_id"])


def badge(label: str, kind: str) -> str:
    return f'<span class="badge {kind}">{escape(label)}</span>'


def go(path: str, msg: str | None = None, error: bool = False):
    if msg:
        sep = "&" if "?" in path else "?"
        path += sep + urlencode({"err" if error else "msg": msg})
    return web.HTTPFound(path)


def to_telegram_html(text: str) -> str:
    """Plain text from the form -> Telegram HTML; **đậm** becomes bold."""
    html = escape(text.strip(), quote=False)
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", html, flags=re.S)


def strip_tags(html: str) -> str:
    return re.sub(r"<[^>]+>", "", html)


def tg(request: web.Request):
    return request.app["tg"]


def database(request: web.Request) -> Database:
    return request.app["tg"].bot_data["db"]


def bot_module():
    # Imported lazily: bot.py imports this module when it starts the web server.
    from . import bot
    return bot


# ---------------------------------------------------------------- layout

CSS = """
:root{--bg:#f4f4f2;--surface:#fcfcfb;--text:#0b0b0b;--text2:#52514e;--muted:#8a8984;--border:#e4e3df;
--accent:#2a78d6;--accent-ink:#fff;--good:#0a7d0a;--good-bg:#e3f4e3;--warn:#8a5a00;--warn-bg:#fdf1d6;
--bad:#c22f2f;--bad-bg:#fbe5e5;--muted-bg:#ecebe8;--grid:#e9e8e4}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#111110;--surface:#1a1a19;--text:#fff;
--text2:#c3c2b7;--muted:#8f8e86;--border:#2f2f2c;--accent:#3987e5;--good:#5fd35f;--good-bg:#173317;
--warn:#f5c35a;--warn-bg:#3a2e10;--bad:#ff7b7b;--bad-bg:#3d1919;--muted-bg:#2a2a28;--grid:#2a2a28}}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--text);font:15px/1.5 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif}
a{color:var(--accent);text-decoration:none}a:hover{text-decoration:underline}
header{background:var(--surface);border-bottom:1px solid var(--border);position:sticky;top:0;z-index:5}
.bar{max-width:1180px;margin:0 auto;padding:0 16px;display:flex;align-items:center;gap:20px;height:56px}
.brand{font-weight:700;white-space:nowrap}
nav{display:flex;gap:4px;overflow-x:auto;flex:1}
nav a{padding:6px 12px;border-radius:8px;color:var(--text2);white-space:nowrap}
nav a.on{background:var(--muted-bg);color:var(--text);font-weight:600}nav a:hover{text-decoration:none;color:var(--text)}
main{max-width:1180px;margin:0 auto;padding:20px 16px 60px}
h1{font-size:22px;margin:0 0 16px}h2{font-size:17px;margin:0 0 12px}
.card{background:var(--surface);border:1px solid var(--border);border-radius:12px;padding:18px;margin-bottom:16px}
.grid{display:grid;gap:16px}.g2{grid-template-columns:repeat(auto-fit,minmax(320px,1fr))}
.tiles{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px;margin-bottom:16px}
.tile{background:var(--surface);border:1px solid var(--border);border-radius:12px;padding:14px 16px}
.tile .k{color:var(--text2);font-size:13px}.tile .v{font-size:24px;font-weight:700;margin-top:2px;font-variant-numeric:tabular-nums}
.tile .s{color:var(--muted);font-size:12px}
.table-wrap{overflow-x:auto}table.wide{min-width:760px}.nowrap{white-space:nowrap}
table{width:100%;border-collapse:collapse;font-variant-numeric:tabular-nums}
th,td{text-align:left;padding:9px 10px;border-bottom:1px solid var(--border);vertical-align:top}
th{color:var(--text2);font-weight:600;font-size:13px;white-space:nowrap}
td.num,th.num{text-align:right}tr:last-child td{border-bottom:0}
.badge{display:inline-block;padding:1px 8px;border-radius:999px;font-size:12px;font-weight:600;white-space:nowrap}
.badge.good{background:var(--good-bg);color:var(--good)}.badge.warn{background:var(--warn-bg);color:var(--warn)}
.badge.bad{background:var(--bad-bg);color:var(--bad)}.badge.muted{background:var(--muted-bg);color:var(--text2)}
label{display:block;font-size:13px;color:var(--text2);margin:0 0 4px}
input,select,textarea{width:100%;padding:8px 10px;border:1px solid var(--border);border-radius:8px;
background:var(--bg);color:var(--text);font:inherit}
textarea{min-height:110px;resize:vertical}
input[type=checkbox],input[type=radio]{width:auto;margin-right:6px}
.field{margin-bottom:12px}.row{display:flex;gap:12px;flex-wrap:wrap;align-items:flex-end}.row>.field{flex:1;min-width:150px}
.check{display:flex;align-items:center;color:var(--text);font-size:14px;margin-bottom:8px}
button,.btn{display:inline-block;padding:8px 14px;border-radius:8px;border:1px solid var(--accent);background:var(--accent);
color:var(--accent-ink);font:inherit;font-weight:600;cursor:pointer;white-space:nowrap}
button:hover,.btn:hover{filter:brightness(1.08);text-decoration:none}
.ghost{background:transparent;color:var(--accent)}.danger{background:transparent;border-color:var(--bad);color:var(--bad)}
.small{padding:4px 10px;font-size:13px}
form.inline{display:inline}
.actions{display:flex;gap:6px;flex-wrap:wrap}
.flash{padding:10px 14px;border-radius:10px;margin-bottom:16px;background:var(--good-bg);color:var(--good);font-weight:600}
.flash.err{background:var(--bad-bg);color:var(--bad)}
.muted{color:var(--muted)}.hint{color:var(--muted);font-size:13px;margin-top:4px}
.mono{font-family:ui-monospace,SFMono-Regular,Menlo,monospace;font-size:13px;word-break:break-all}
.chart svg{width:100%;height:auto;display:block}
.chart .col .hit{fill:transparent}.chart .col:hover .b{fill:var(--text)}
.chart .b{fill:var(--accent)}.chart .gl{stroke:var(--grid);stroke-width:1}
.chart text{fill:var(--muted);font-size:11px}
details summary{cursor:pointer;color:var(--accent);margin-top:10px}
.pager{display:flex;gap:8px;align-items:center;margin-top:12px}
.thumb{width:56px;height:56px;object-fit:cover;border-radius:6px;border:1px solid var(--border)}
.login{max-width:360px;margin:12vh auto}
.chart .m{display:none}
@media (max-width:640px){.bar{flex-wrap:wrap;height:auto;padding:10px 16px;gap:8px 12px}
.brand{flex:1}nav{order:3;flex-basis:100%}.chart .d{display:none}.chart .m{display:block}}
"""


def page(request: web.Request, title: str, body: str, active: str = "") -> web.Response:
    nav = "".join(
        f'<a href="{href}" class="{"on" if key == active else ""}">{label}</a>'
        for key, href, label in [
            ("home", "/admin", "Tổng quan"),
            ("products", "/admin/products", "Sản phẩm & kho"),
            ("orders", "/admin/orders", "Đơn hàng"),
            ("posts", "/admin/posts", "Thông báo"),
        ]
    )
    flash = ""
    if request.query.get("msg"):
        flash = f'<div class="flash">{escape(request.query["msg"])}</div>'
    elif request.query.get("err"):
        flash = f'<div class="flash err">{escape(request.query["err"])}</div>'
    html = f"""<!doctype html><html lang="vi"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{escape(title)} · Quản trị shop</title>
<style>{CSS}</style></head><body>
<header><div class="bar"><span class="brand">🛍 Quản trị shop</span><nav>{nav}</nav>
<a class="btn ghost small" href="/admin/backup.zip" title="Tải bản sao lưu dữ liệu">💾 Sao lưu</a>
<form class="inline" method="post" action="/admin/logout"><button class="ghost small">Đăng xuất</button></form></div></header>
<main>{flash}{body}</main></body></html>"""
    return web.Response(text=html, content_type="text/html")


# ---------------------------------------------------------------- auth

@web.middleware
async def require_login(request: web.Request, handler):
    if request.path.startswith("/admin") and request.path != "/admin/login":
        if request.cookies.get(COOKIE) not in request.app["sessions"]:
            raise web.HTTPFound("/admin/login")
    return await handler(request)


async def login_form(request: web.Request) -> web.Response:
    error = '<div class="flash err">Sai mật khẩu.</div>' if request.query.get("fail") else ""
    html = f"""<!doctype html><html lang="vi"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Đăng nhập · Quản trị shop</title>
<style>{CSS}</style></head><body><main><div class="login card"><h1>🛍 Quản trị shop</h1>{error}
<form method="post"><div class="field"><label>Mật khẩu</label>
<input type="password" name="password" autofocus required></div><button>Đăng nhập</button></form></div></main></body></html>"""
    return web.Response(text=html, content_type="text/html")


async def login(request: web.Request) -> web.Response:
    form = await request.post()
    expected = tg(request).bot_data["config"].admin_password.encode()
    if not hmac.compare_digest(str(form.get("password", "")).encode(), expected):
        log.warning("Failed admin login from %s", request.remote)
        await asyncio.sleep(1.5)  # slow down password guessing
        raise web.HTTPFound("/admin/login?fail=1")
    token = secrets.token_urlsafe(32)
    request.app["sessions"].add(token)
    response = web.HTTPFound("/admin")
    response.set_cookie(COOKIE, token, max_age=30 * 86400, httponly=True, samesite="Strict")
    raise response


async def logout(request: web.Request) -> web.Response:
    request.app["sessions"].discard(request.cookies.get(COOKIE))
    response = web.HTTPFound("/admin/login")
    response.del_cookie(COOKIE)
    raise response


# ---------------------------------------------------------------- overview

def nice_max(value: int) -> int:
    if value <= 0:
        return 100_000
    magnitude = 10 ** (len(str(int(value))) - 1)
    for step in (1, 2, 2.5, 5, 10):
        if value <= step * magnitude:
            return int(step * magnitude)
    return value


def short_money(amount: int) -> str:
    if amount >= 1_000_000:
        return f"{amount / 1_000_000:g}tr"
    if amount >= 1_000:
        return f"{amount / 1_000:g}k"
    return str(amount)


def revenue_chart(days: list[tuple[date, int, int]], width: int = 720, label_every: int = 5) -> str:
    """Single-series bar chart: revenue per day, native tooltip per column."""
    height, left, bottom, top = 220, 48, 22, 10
    plot_w, plot_h = width - left - 8, height - top - bottom
    peak = nice_max(max((r for _, r, _ in days), default=0))
    step = plot_w / len(days)
    bar_w = max(step - 4, 2)  # 4px gap between bars
    base = top + plot_h
    parts = []
    for frac in (0, 0.5, 1):
        y = base - plot_h * frac
        parts.append(f'<line class="gl" x1="{left}" x2="{width - 8}" y1="{y:.1f}" y2="{y:.1f}"/>')
        parts.append(f'<text x="{left - 6}" y="{y + 4:.1f}" text-anchor="end">{short_money(int(peak * frac))}</text>')
    for i, (day, revenue, orders) in enumerate(days):
        x = left + i * step + (step - bar_w) / 2
        h = plot_h * revenue / peak if revenue else 0
        tip = f"{day:%d/%m}: {money(revenue)} · {orders} đơn"
        bar = ""
        if h > 0:
            r = min(4, h, bar_w / 2)  # rounded data end, square at the baseline
            y = base - h
            bar = (f'<path class="b" d="M{x:.1f},{base} V{y + r:.1f} Q{x:.1f},{y:.1f} {x + r:.1f},{y:.1f} '
                   f'H{x + bar_w - r:.1f} Q{x + bar_w:.1f},{y:.1f} {x + bar_w:.1f},{y + r:.1f} V{base} Z"/>')
        parts.append(
            f'<g class="col"><title>{tip}</title>'
            f'<rect class="hit" x="{left + i * step:.1f}" y="{top}" width="{step:.1f}" height="{plot_h}"/>{bar}</g>'
        )
        if (len(days) - 1 - i) % label_every == 0:
            parts.append(f'<text x="{x + bar_w / 2:.1f}" y="{height - 6}" text-anchor="middle">{day:%d/%m}</text>')
    return f'<svg viewBox="0 0 {width} {height}" role="img" aria-label="Doanh thu {len(days)} ngày">{"".join(parts)}</svg>'


async def overview(request: web.Request) -> web.Response:
    db = database(request)
    s = db.revenue_summary()
    pending = db.search_orders(status="pending", limit=0)[1]
    by_day = db.revenue_by_day(CHART_DAYS)
    today = date.today()
    days = []
    for offset in range(CHART_DAYS - 1, -1, -1):
        day = today - timedelta(days=offset)
        revenue, orders = by_day.get(day.isoformat(), (0, 0))
        days.append((day, revenue, orders))

    tiles = "".join(
        f'<div class="tile"><div class="k">{k}</div><div class="v">{v}</div><div class="s">{sub}</div></div>'
        for k, v, sub in [
            ("Doanh thu hôm nay", money(s["today"]), f"{s['today_orders']} đơn"),
            ("7 ngày qua", money(s["week"]), ""),
            ("Tháng này", money(s["month"]), ""),
            ("Tổng doanh thu", money(s["total"]), f"{s['total_orders']} đơn đã giao"),
            ("Đơn chờ thanh toán", str(pending), '<a href="/admin/orders?status=pending">Xem đơn</a>'),
            ("Người dùng bot", str(db.count_users()), "nhận được thông báo"),
        ]
    )
    day_rows = "".join(
        f"<tr><td>{d:%d/%m/%Y}</td><td class='num'>{o}</td><td class='num'>{money(r)}</td></tr>"
        for d, r, o in reversed(days) if o
    ) or "<tr><td colspan=3 class='muted'>Chưa có đơn nào.</td></tr>"

    top = db.top_products(CHART_DAYS)
    top_rows = "".join(
        f"<tr><td>{escape(p['name'])}</td><td class='num'>{p['sold']}</td><td class='num'>{money(p['revenue'])}</td></tr>"
        for p in top
    ) or "<tr><td colspan=3 class='muted'>Chưa có dữ liệu.</td></tr>"

    low = [
        (p, db.sellable_quantity(p["id"])) for p in db.list_products(only_active=True)
    ]
    low = [(p, n) for p, n in low if n <= LOW_STOCK]
    low_rows = "".join(
        f"<tr><td><a href='/admin/products/{p['id']}'>{escape(p['name'])}</a></td>"
        f"<td class='num'>{badge(str(n), 'bad' if n == 0 else 'warn')}</td></tr>"
        for p, n in low
    ) or "<tr><td colspan=2 class='muted'>Mọi sản phẩm đều đủ hàng.</td></tr>"

    body = f"""<h1>Tổng quan</h1><div class="tiles">{tiles}</div>
<div class="card chart"><h2>Doanh thu theo ngày</h2>
<div class="d">{revenue_chart(days)}<p class="hint">{CHART_DAYS} ngày qua · rê chuột vào cột để xem số</p></div>
<div class="m">{revenue_chart(days[-14:], width=360, label_every=3)}<p class="hint">14 ngày qua · chạm vào cột để xem số</p></div>
<details><summary>Xem dạng bảng</summary><div class="table-wrap"><table>
<tr><th>Ngày</th><th class="num">Đơn</th><th class="num">Doanh thu</th></tr>{day_rows}</table></div></details></div>
<div class="grid g2">
<div class="card"><h2>Bán chạy {CHART_DAYS} ngày</h2><div class="table-wrap"><table>
<tr><th>Sản phẩm</th><th class="num">Đã bán</th><th class="num">Doanh thu</th></tr>{top_rows}</table></div></div>
<div class="card"><h2>Sắp hết hàng (≤ {LOW_STOCK})</h2><div class="table-wrap"><table>
<tr><th>Sản phẩm</th><th class="num">Còn bán</th></tr>{low_rows}</table></div></div></div>"""
    return page(request, "Tổng quan", body, "home")


# ---------------------------------------------------------------- products & stock

async def products(request: web.Request) -> web.Response:
    db = database(request)
    rows = []
    for p in db.list_products(only_active=False):
        sellable = db.sellable_quantity(p["id"])
        reserved = db.reserved_quantity(p["id"])
        state = badge("Đang bán", "good") if p["active"] else badge("Đang ẩn", "muted")
        rows.append(
            f"<tr><td class='num'>{p['id']}</td><td><a href='/admin/products/{p['id']}'>{escape(p['name'])}</a></td>"
            f"<td class='num'>{money(p['price'])}</td>"
            f"<td class='num'>{badge(str(sellable), 'bad' if sellable == 0 else 'warn' if sellable <= LOW_STOCK else 'good')}</td>"
            f"<td class='num'>{reserved}</td><td class='num'>{db.delivered_count(p['id'])}</td><td>{state}</td>"
            f"<td><a class='btn small ghost' href='/admin/products/{p['id']}'>Sửa / nhập kho</a></td></tr>"
        )
    table = "".join(rows) or "<tr><td colspan=8 class='muted'>Chưa có sản phẩm.</td></tr>"
    body = f"""<h1>Sản phẩm & kho</h1>
<div class="card"><h2>Thêm sản phẩm</h2><form method="post" action="/admin/products">
<div class="row"><div class="field" style="flex:2"><label>Tên sản phẩm</label><input name="name" required></div>
<div class="field"><label>Giá (đồng)</label><input name="price" inputmode="numeric" placeholder="65000" required></div></div>
<div class="field"><label>Mô tả</label><textarea name="description" style="min-height:70px"></textarea></div>
<button>Thêm sản phẩm</button></form></div>
<div class="card"><div class="table-wrap"><table class="wide">
<tr><th class="num">ID</th><th>Tên</th><th class="num">Giá</th><th class="num">Còn bán</th><th class="num">Đang giữ</th>
<th class="num">Đã bán</th><th>Trạng thái</th><th></th></tr>{table}</table></div>
<p class="hint">"Đang giữ" là hàng đã được giữ cho đơn chờ thanh toán.</p></div>"""
    return page(request, "Sản phẩm", body, "products")


async def create_product(request: web.Request) -> web.Response:
    form = await request.post()
    try:
        name = str(form.get("name", "")).strip()
        if not name:
            raise ValueError("Thiếu tên sản phẩm")
        product_id = database(request).add_product(name, parse_price(str(form.get("price"))), str(form.get("description", "")).strip())
    except ValueError as e:
        raise go("/admin/products", str(e), error=True)
    raise go(f"/admin/products/{product_id}", "Đã thêm sản phẩm. Giờ hãy nhập kho bên dưới.")


def get_product_or_404(request: web.Request):
    product = database(request).get_product(int(request.match_info["id"]))
    if product is None:
        raise web.HTTPNotFound(text="Không tìm thấy sản phẩm")
    return product


async def product_detail(request: web.Request) -> web.Response:
    db = database(request)
    p = get_product_or_404(request)
    sellable = db.sellable_quantity(p["id"])
    available = db.list_stock(p["id"], available=True)
    delivered = db.list_stock(p["id"], available=False, limit=100)
    avail_rows = "".join(
        f"<tr><td class='mono'>{escape(s['content'])}</td><td><form class='inline' method='post' "
        f"action='/admin/stock/{s['id']}/delete' onsubmit=\"return confirm('Xoá hàng này khỏi kho?')\">"
        f"<button class='danger small'>Xoá</button></form></td></tr>"
        for s in available
    ) or "<tr><td colspan=2 class='muted'>Kho trống.</td></tr>"
    sold_rows = "".join(
        f"<tr><td class='mono'>{escape(s['content'])}</td><td><a href='/admin/orders/{s['order_id']}'>{s['code']}</a></td></tr>"
        for s in delivered
    ) or "<tr><td colspan=2 class='muted'>Chưa bán hàng nào.</td></tr>"
    toggle_label = "Ẩn sản phẩm" if p["active"] else "Hiện sản phẩm"
    auto = "đang BẬT" if bot_module().auto_restock_enabled(tg(request)) else "đang TẮT"
    notify_btn = (
        f"<form class='inline' method='post' action='/admin/products/{p['id']}/notify'>"
        f"<button class='ghost'>📣 Gửi thông báo có hàng</button></form>" if sellable and p["active"] else ""
    )
    body = f"""<p><a href="/admin/products">← Tất cả sản phẩm</a></p>
<h1>#{p['id']} {escape(p['name'])} {badge('Đang bán', 'good') if p['active'] else badge('Đang ẩn', 'muted')}</h1>
<div class="tiles"><div class="tile"><div class="k">Còn bán</div><div class="v">{sellable}</div></div>
<div class="tile"><div class="k">Đang giữ cho đơn chờ</div><div class="v">{db.reserved_quantity(p['id'])}</div></div>
<div class="tile"><div class="k">Đã bán</div><div class="v">{db.delivered_count(p['id'])}</div></div></div>
<div class="grid g2">
<div class="card"><h2>Thông tin</h2><form method="post" action="/admin/products/{p['id']}">
<div class="field"><label>Tên</label><input name="name" value="{escape(p['name'])}" required></div>
<div class="field"><label>Giá (đồng)</label><input name="price" value="{p['price']}" inputmode="numeric" required></div>
<div class="field"><label>Mô tả</label><textarea name="description">{escape(p['description'])}</textarea></div>
<button>Lưu</button></form>
<div class="actions" style="margin-top:12px">
<form class="inline" method="post" action="/admin/products/{p['id']}/toggle"><button class="ghost">{toggle_label}</button></form>
{notify_btn}</div></div>
<div class="card"><h2>Nhập kho</h2>
<form method="post" action="/admin/products/{p['id']}/stock" enctype="multipart/form-data">
<div class="field"><label>File Excel (.xlsx), CSV hoặc TXT</label><input type="file" name="file" accept=".xlsx,.csv,.txt">
<p class="hint">Mỗi dòng là 1 hàng. Nhiều cột sẽ được ghép: <span class="mono">email | matkhau | 2fa</span></p></div>
<label class="check"><input type="checkbox" name="header" checked>Dòng đầu của file là tiêu đề (bỏ qua)</label>
<div class="field"><label>Hoặc dán trực tiếp (mỗi dòng 1 hàng)</label><textarea name="text" placeholder="user1@gmail.com | matkhau1"></textarea></div>
<label class="check"><input type="checkbox" name="dedupe" checked>Bỏ qua hàng trùng với hàng đã có</label>
<button>Nhập kho</button>
<p class="hint">Thông báo "có hàng lại" tự động {auto} — đổi ở trang <a href="/admin/posts">Thông báo</a>.</p>
</form></div></div>
<div class="grid g2">
<div class="card"><h2>Hàng trong kho ({len(available)})</h2><div class="table-wrap"><table>{avail_rows}</table></div></div>
<div class="card"><h2>Đã giao gần đây</h2><div class="table-wrap"><table>
<tr><th>Hàng</th><th>Đơn</th></tr>{sold_rows}</table></div></div></div>"""
    return page(request, p["name"], body, "products")


async def update_product(request: web.Request) -> web.Response:
    p = get_product_or_404(request)
    form = await request.post()
    try:
        name = str(form.get("name", "")).strip()
        if not name:
            raise ValueError("Thiếu tên sản phẩm")
        database(request).update_product(p["id"], name, parse_price(str(form.get("price"))), str(form.get("description", "")).strip())
    except ValueError as e:
        raise go(f"/admin/products/{p['id']}", str(e), error=True)
    raise go(f"/admin/products/{p['id']}", "Đã lưu.")


async def toggle_product(request: web.Request) -> web.Response:
    p = get_product_or_404(request)
    database(request).set_product_active(p["id"], not p["active"])
    raise go(f"/admin/products/{p['id']}", "Đã ẩn sản phẩm." if p["active"] else "Sản phẩm đã hiện lại.")


async def notify_product(request: web.Request) -> web.Response:
    p = get_product_or_404(request)
    app = tg(request)
    app.create_task(bot_module().announce_product(app, p["id"]))
    raise go(f"/admin/products/{p['id']}", "Đang gửi thông báo có hàng cho khách. Bot sẽ báo kết quả qua Telegram.")


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores 123456 as 123456.0
    return str(value).strip()


def parse_stock_file(filename: str, data: bytes, skip_header: bool) -> list[str]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [list(r) for r in workbook.active.iter_rows(values_only=True)]
        workbook.close()
    elif name.endswith(".csv"):
        text = decode_text(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    elif name.endswith(".txt"):
        rows = [[line] for line in decode_text(data).splitlines()]
    else:
        raise ValueError("Chỉ nhận file .xlsx, .csv hoặc .txt (file .xls cũ hãy mở bằng Excel rồi lưu lại dạng .xlsx).")
    if skip_header and rows:
        rows = rows[1:]
    items = [" | ".join(t for t in (cell_text(c) for c in row) if t) for row in rows]
    return [i for i in items if i]


async def import_stock(request: web.Request) -> web.Response:
    db = database(request)
    p = get_product_or_404(request)
    back = f"/admin/products/{p['id']}"
    form = await request.post()
    items: list[str] = []
    upload = form.get("file")
    try:
        if upload is not None and getattr(upload, "filename", ""):
            items += parse_stock_file(upload.filename, upload.file.read(), bool(form.get("header")))
    except ValueError as e:
        raise go(back, str(e), error=True)
    except Exception:
        log.exception("Could not read stock file")
        raise go(back, "Không đọc được file. Hãy kiểm tra lại định dạng.", error=True)
    items += [line.strip() for line in str(form.get("text", "")).splitlines() if line.strip()]
    if not items:
        raise go(back, "Không có hàng nào để nhập.", error=True)

    skipped = 0
    if form.get("dedupe"):
        seen = db.stock_contents(p["id"])
        unique = []
        for item in items:
            if item in seen:
                skipped += 1
            else:
                seen.add(item)
                unique.append(item)
        items = unique
    if not items:
        raise go(back, f"Tất cả {skipped} hàng đều đã có trong kho, không nhập thêm.", error=True)

    was_sold_out = db.sellable_quantity(p["id"]) == 0
    db.add_stock(p["id"], items)
    msg = f"Đã nhập {len(items)} hàng." + (f" Bỏ qua {skipped} hàng trùng." if skipped else "")
    if bot_module().after_restock(tg(request), p["id"], was_sold_out):
        msg += " Đang gửi thông báo có hàng cho khách."
    raise go(back, msg)


async def delete_stock(request: web.Request) -> web.Response:
    product_id = database(request).delete_stock_item(int(request.match_info["id"]))
    if product_id is None:
        raise go("/admin/products", "Không xoá được (hàng đã bán hoặc không còn).", error=True)
    raise go(f"/admin/products/{product_id}", "Đã xoá 1 hàng khỏi kho.")


# ---------------------------------------------------------------- orders

def order_filters(request: web.Request) -> dict:
    q = request.query
    status = q.get("status", "")
    return {
        "status": status if status in STATUS else None,
        "query": q.get("q", "").strip() or None,
        "date_from": q.get("from") or None,
        "date_to": q.get("to") or None,
    }


def order_actions(o, back: str) -> str:
    if o["status"] != "pending":
        return ""
    return (
        f"<div class='actions'><form class='inline' method='post' action='/admin/orders/{o['id']}/confirm' "
        f"onsubmit=\"return confirm('Xác nhận đã nhận {money(o['amount'])} cho đơn {o['code']} và giao hàng?')\">"
        f"<input type='hidden' name='back' value='{escape(back)}'><button class='small'>✅ Đã nhận tiền</button></form>"
        f"<form class='inline' method='post' action='/admin/orders/{o['id']}/cancel' "
        f"onsubmit=\"return confirm('Huỷ đơn {o['code']}?')\"><input type='hidden' name='back' value='{escape(back)}'>"
        f"<button class='danger small'>Huỷ</button></form></div>"
    )


async def orders(request: web.Request) -> web.Response:
    db = database(request)
    f = order_filters(request)
    page_no = int(request.query["page"]) if request.query.get("page", "").isdigit() else 1
    page_no = max(page_no, 1)
    rows, total = db.search_orders(**f, limit=PAGE_SIZE, offset=(page_no - 1) * PAGE_SIZE)
    back = request.path_qs
    table = "".join(
        f"<tr><td class='nowrap'><a href='/admin/orders/{o['id']}' class='mono'>{o['code']}</a></td>"
        f"<td class='nowrap'>{local_time(o['created_at'])}</td><td>{escape(customer(o))}</td>"
        f"<td>{escape(o['product_name'])}</td><td class='num'>{o['quantity']}</td>"
        f"<td class='num'>{money(o['amount'])}</td><td>{badge(*STATUS[o['status']])}</td>"
        f"<td>{order_actions(o, back)}</td></tr>"
        for o in rows
    ) or "<tr><td colspan=8 class='muted'>Không có đơn phù hợp.</td></tr>"
    options = "".join(
        f"<option value='{k}' {'selected' if f['status'] == k else ''}>{v[0]}</option>" for k, v in STATUS.items()
    )
    query = {k: v for k, v in request.query.items() if k in ("status", "q", "from", "to") and v}
    pages = max((total + PAGE_SIZE - 1) // PAGE_SIZE, 1)
    pager = ""
    if pages > 1:
        prev_link = f"<a class='btn small ghost' href='?{urlencode({**query, 'page': page_no - 1})}'>← Trước</a>" if page_no > 1 else ""
        next_link = f"<a class='btn small ghost' href='?{urlencode({**query, 'page': page_no + 1})}'>Sau →</a>" if page_no < pages else ""
        pager = f"<div class='pager'>{prev_link}<span class='muted'>Trang {page_no}/{pages}</span>{next_link}</div>"
    body = f"""<h1>Đơn hàng</h1>
<div class="card"><form method="get" class="row">
<div class="field"><label>Trạng thái</label><select name="status"><option value="">Tất cả</option>{options}</select></div>
<div class="field" style="flex:2"><label>Tìm (mã đơn, @khách, sản phẩm)</label><input name="q" value="{escape(request.query.get('q', ''))}"></div>
<div class="field"><label>Từ ngày</label><input type="date" name="from" value="{escape(request.query.get('from', ''))}"></div>
<div class="field"><label>Đến ngày</label><input type="date" name="to" value="{escape(request.query.get('to', ''))}"></div>
<div class="field actions" style="flex:0"><button>Lọc</button>
<a class="btn ghost" href="/admin/orders.xlsx?{urlencode(query)}">⬇ Xuất Excel</a></div></form></div>
<div class="card"><p class="muted" style="margin-top:0">{total} đơn</p><div class="table-wrap"><table class="wide">
<tr><th>Mã</th><th>Thời gian</th><th>Khách</th><th>Sản phẩm</th><th class="num">SL</th><th class="num">Tiền</th>
<th>Trạng thái</th><th></th></tr>{table}</table></div>{pager}</div>"""
    return page(request, "Đơn hàng", body, "orders")


def get_order_or_404(request: web.Request):
    order = database(request).get_order(int(request.match_info["id"]))
    if order is None:
        raise web.HTTPNotFound(text="Không tìm thấy đơn")
    return order


async def order_detail(request: web.Request) -> web.Response:
    o = get_order_or_404(request)
    items = database(request).delivered_items(o["id"])
    item_rows = "".join(f"<tr><td class='mono'>{escape(i)}</td></tr>" for i in items)
    resend = (
        f"<form method='post' action='/admin/orders/{o['id']}/resend'><button class='ghost'>📨 Gửi lại hàng cho khách</button></form>"
        if o["status"] == "paid" else ""
    )
    body = f"""<p><a href="/admin/orders">← Tất cả đơn</a></p>
<h1>Đơn <span class="mono" style="font-size:20px">{o['code']}</span> {badge(*STATUS[o['status']])}</h1>
<div class="grid g2"><div class="card"><table>
<tr><th>Sản phẩm</th><td>{escape(o['product_name'])} × {o['quantity']}</td></tr>
<tr><th>Số tiền</th><td>{money(o['amount'])}</td></tr>
<tr><th>Khách</th><td>{escape(customer(o))} <span class="muted">(ID {o['user_id']})</span></td></tr>
<tr><th>Tạo lúc</th><td>{local_time(o['created_at'])}</td></tr>
<tr><th>Thanh toán lúc</th><td>{local_time(o['paid_at']) or '—'}</td></tr></table>
<div style="margin-top:12px">{order_actions(o, request.path)}{resend}</div></div>
<div class="card"><h2>Hàng đã giao</h2><div class="table-wrap"><table>
{item_rows or "<tr><td class='muted'>Chưa giao.</td></tr>"}</table></div></div></div>"""
    return page(request, o["code"], body, "orders")


def safe_back(value, default: str) -> str:
    value = str(value or "")
    return value if value.startswith("/admin") else default


async def confirm_order(request: web.Request) -> web.Response:
    o = get_order_or_404(request)
    back = safe_back((await request.post()).get("back"), f"/admin/orders/{o['id']}")
    if await bot_module().deliver(tg(request), o["id"]):
        raise go(back, f"Đã giao hàng cho đơn {o['code']}.")
    raise go(back, f"Không giao được đơn {o['code']}: đơn không còn chờ thanh toán hoặc kho không đủ hàng.", error=True)


async def cancel_order(request: web.Request) -> web.Response:
    o = get_order_or_404(request)
    back = safe_back((await request.post()).get("back"), f"/admin/orders/{o['id']}")
    if not database(request).cancel_order(o["id"]):
        raise go(back, "Không huỷ được đơn này.", error=True)
    try:
        await tg(request).bot.send_message(o["user_id"], f"❌ Đơn {o['code']} đã bị huỷ bởi cửa hàng.")
    except TelegramError:
        pass
    raise go(back, f"Đã huỷ đơn {o['code']}.")


async def resend_order(request: web.Request) -> web.Response:
    o = get_order_or_404(request)
    try:
        await bot_module().send_items(tg(request), o, database(request).delivered_items(o["id"]))
    except TelegramError as e:
        raise go(f"/admin/orders/{o['id']}", f"Không gửi được: {e}", error=True)
    raise go(f"/admin/orders/{o['id']}", "Đã gửi lại hàng cho khách.")


async def export_orders(request: web.Request) -> web.Response:
    rows, _ = database(request).search_orders(**order_filters(request), limit=None)
    wb = Workbook()
    ws = wb.active
    ws.title = "Đơn hàng"
    ws.append(["Mã đơn", "Tạo lúc", "Thanh toán lúc", "Khách", "ID khách", "Sản phẩm", "Số lượng", "Số tiền", "Trạng thái"])
    per_day: dict[str, list[int]] = {}
    for o in rows:
        ws.append([
            o["code"], local_time(o["created_at"]), local_time(o["paid_at"]), customer(o), o["user_id"],
            o["product_name"], o["quantity"], o["amount"], STATUS[o["status"]][0],
        ])
        if o["status"] == "paid":
            day = local_time(o["paid_at"], "%Y-%m-%d")
            per_day.setdefault(day, [0, 0])
            per_day[day][0] += 1
            per_day[day][1] += o["amount"]
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGHI", (12, 17, 17, 18, 13, 30, 10, 13, 15)):
        ws.column_dimensions[col].width = width
    for cell in ws["H"][1:]:
        cell.number_format = "#,##0"

    ws2 = wb.create_sheet("Doanh thu theo ngày")
    ws2.append(["Ngày", "Số đơn", "Doanh thu"])
    for day in sorted(per_day):
        ws2.append([datetime.strptime(day, "%Y-%m-%d").strftime("%d/%m/%Y"), *per_day[day]])
    ws2.append(["Tổng", sum(v[0] for v in per_day.values()), sum(v[1] for v in per_day.values())])
    for cell in ws2[1] + ws2[ws2.max_row]:
        cell.font = Font(bold=True)
    for cell in ws2["C"][1:]:
        cell.number_format = "#,##0"
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["C"].width = 15

    buffer = io.BytesIO()
    wb.save(buffer)
    filename = f"don-hang-{date.today():%Y%m%d}.xlsx"
    return web.Response(
        body=buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------- promotional posts

async def posts(request: web.Request) -> web.Response:
    db = database(request)
    app = tg(request)
    product_options = "".join(
        f"<option value='{p['id']}'>#{p['id']} {escape(p['name'])}</option>" for p in db.list_products()
    )
    default_time = (datetime.now() + timedelta(hours=1)).replace(minute=0).strftime("%Y-%m-%dT%H:%M")
    rows = []
    for po in db.list_posts():
        label, kind = POST_STATUS.get(po["status"], (po["status"], "muted"))
        result = f"{po['delivered']} nhận · {po['failed']} lỗi" if po["status"] == "sent" else ""
        thumb = f"<img class='thumb' src='/admin/uploads/{escape(po['image'])}' alt=''>" if po["image"] else ""
        cancel = (
            f"<form class='inline' method='post' action='/admin/posts/{po['id']}/cancel' "
            f"onsubmit=\"return confirm('Huỷ tin đã hẹn giờ?')\"><button class='danger small'>Huỷ</button></form>"
            if po["status"] == "scheduled" else ""
        )
        excerpt = strip_tags(po["text"])
        excerpt = excerpt[:140] + ("…" if len(excerpt) > 140 else "")
        rows.append(
            f"<tr><td>{local_time(po['send_at'])}</td><td>{thumb}</td><td>{escape(excerpt)}</td>"
            f"<td>{escape(po['product_name'] or '—')}</td><td>{badge(label, kind)}<div class='hint'>{result}</div></td>"
            f"<td>{cancel}</td></tr>"
        )
    table = "".join(rows) or "<tr><td colspan=6 class='muted'>Chưa có tin nào.</td></tr>"
    auto_on = bot_module().auto_restock_enabled(app)
    body = f"""<h1>Thông báo</h1>
<div class="grid g2">
<div class="card"><h2>Soạn tin khuyến mãi</h2>
<form id="post-form" method="post" action="/admin/posts" enctype="multipart/form-data">
<div class="field"><label>Nội dung</label><textarea name="text" required
placeholder="🔥 KHUYẾN MÃI Capcut 1 tháng chỉ **30K**!"></textarea>
<p class="hint">Viết **chữ** để in đậm. Có ảnh thì tối đa 1024 ký tự, không ảnh tối đa 4096.</p></div>
<div class="field"><label>Ảnh banner (không bắt buộc)</label><input type="file" name="image" accept=".jpg,.jpeg,.png,.webp"></div>
<div class="field"><label>Nút 🛒 Mua ngay cho sản phẩm</label><select name="product_id">
<option value="">Không gắn nút</option>{product_options}</select></div>
<label class="check"><input type="radio" name="when" value="now" checked>Gửi ngay</label>
<div class="row"><label class="check" style="margin:0 0 12px"><input type="radio" name="when" value="later">Hẹn giờ gửi lúc</label>
<div class="field"><input type="datetime-local" name="send_at" value="{default_time}"></div></div>
<div class="actions"><button type="button" class="ghost" id="test-btn">Gửi thử cho tôi</button>
<button onclick="return confirm('Gửi / hẹn giờ gửi tin này cho {db.count_users()} khách?')">Gửi cho tất cả khách</button></div>
<p class="hint" id="test-status"></p></form></div>
<div class="card"><h2>Thông báo tự động</h2>
<p>"🔥 TIN NÓNG … đã có hàng lại!" được gửi cho mọi khách khi bạn nhập kho cho sản phẩm đang hết hàng.</p>
<p>Trạng thái: {badge('Đang bật', 'good') if auto_on else badge('Đang tắt', 'muted')}</p>
<form method="post" action="/admin/settings/auto-restock"><input type="hidden" name="on" value="{'0' if auto_on else '1'}">
<button class="{'danger' if auto_on else ''}">{'Tắt' if auto_on else 'Bật'} thông báo có hàng</button></form>
<p class="hint">Đổi biểu tượng (kể cả emoji động) bằng lệnh /setemoji trong bot.</p></div></div>
<div class="card"><h2>Tin đã gửi & đã hẹn giờ</h2><div class="table-wrap"><table class="wide">
<tr><th>Thời gian</th><th>Ảnh</th><th>Nội dung</th><th>Nút Mua ngay</th><th>Trạng thái</th><th></th></tr>{table}</table></div></div>
<script>
document.getElementById('test-btn').addEventListener('click', async (e) => {{
  const btn = e.currentTarget, status = document.getElementById('test-status');
  btn.disabled = true; status.textContent = 'Đang gửi thử…';
  try {{
    const r = await fetch('/admin/posts/test', {{method: 'POST', body: new FormData(document.getElementById('post-form'))}});
    status.textContent = (await r.json()).message;
  }} catch (err) {{ status.textContent = 'Lỗi: ' + err; }}
  btn.disabled = false;
}});
</script>"""
    return page(request, "Thông báo", body, "posts")


async def read_post_form(request: web.Request) -> dict:
    """Validate the compose form; saves the image. Raises ValueError with a message."""
    form = await request.post()
    raw = str(form.get("text", "")).strip()
    if not raw:
        raise ValueError("Hãy nhập nội dung tin.")
    text = to_telegram_html(raw)
    image = None
    upload = form.get("image")
    if upload is not None and getattr(upload, "filename", ""):
        ext = "." + upload.filename.rsplit(".", 1)[-1].lower() if "." in upload.filename else ""
        if ext not in IMAGE_TYPES:
            raise ValueError("Ảnh phải là .jpg, .png hoặc .webp.")
        data = upload.file.read()
        if len(data) > 10 * 1024 * 1024:
            raise ValueError("Ảnh quá lớn (tối đa 10MB).")
        image = secrets.token_hex(8) + ext
        (bot_module().uploads_dir(tg(request)) / image).write_bytes(data)
    limit = 1024 if image else 4096
    if len(strip_tags(text)) > limit:
        raise ValueError(f"Nội dung quá dài ({len(strip_tags(text))} ký tự, tối đa {limit}).")
    product_id = int(form["product_id"]) if form.get("product_id") else None
    send_at = datetime.now(timezone.utc)
    if form.get("when") == "later":
        try:
            send_at = datetime.fromisoformat(str(form.get("send_at"))).astimezone()  # local time from the browser
        except ValueError:
            raise ValueError("Thời gian hẹn không hợp lệ.")
        if send_at < datetime.now(timezone.utc) - timedelta(minutes=1):
            raise ValueError("Thời gian hẹn đã qua.")
    return {"text": text, "image": image, "product_id": product_id, "send_at": send_at}


async def test_post(request: web.Request) -> web.Response:
    try:
        post = await read_post_form(request)
    except ValueError as e:
        return web.json_response({"message": f"⚠️ {e}"})
    admins = list(tg(request).bot_data["config"].admin_ids)
    try:
        await bot_module().send_post(tg(request), post, chat_ids=admins)
    except TelegramError as e:
        return web.json_response({"message": f"⚠️ Telegram báo lỗi: {e}"})
    finally:
        if post["image"]:
            (bot_module().uploads_dir(tg(request)) / post["image"]).unlink(missing_ok=True)
    return web.json_response({"message": "✅ Đã gửi thử vào Telegram của bạn. Kiểm tra rồi bấm \"Gửi cho tất cả khách\"."})


async def create_post(request: web.Request) -> web.Response:
    try:
        post = await read_post_form(request)
    except ValueError as e:
        raise go("/admin/posts", str(e), error=True)
    database(request).create_post(**post)
    app = tg(request)
    if post["send_at"] <= datetime.now(timezone.utc):
        app.create_task(bot_module().run_due_posts(app))
        raise go("/admin/posts", "Đang gửi tin cho khách. Bot sẽ báo kết quả qua Telegram.")
    raise go("/admin/posts", f"Đã hẹn gửi lúc {post['send_at'].astimezone():%H:%M %d/%m/%Y}.")


async def cancel_post(request: web.Request) -> web.Response:
    ok = database(request).cancel_post(int(request.match_info["id"]))
    raise go("/admin/posts", "Đã huỷ tin hẹn giờ." if ok else "Tin này không còn ở trạng thái hẹn giờ.", error=not ok)


async def upload_file(request: web.Request) -> web.Response:
    name = request.match_info["name"]
    if not re.fullmatch(r"[0-9a-f]{16}\.(jpg|jpeg|png|webp)", name):
        raise web.HTTPNotFound()
    path = bot_module().uploads_dir(tg(request)) / name
    if not path.exists():
        raise web.HTTPNotFound()
    return web.FileResponse(path)


async def set_auto_restock(request: web.Request) -> web.Response:
    on = (await request.post()).get("on") == "1"
    database(request).set_setting("auto_restock", "1" if on else "0")
    raise go("/admin/posts", "Đã bật thông báo có hàng tự động." if on else "Đã tắt thông báo có hàng tự động.")


# ---------------------------------------------------------------- wiring

async def download_backup(request: web.Request) -> web.Response:
    path = bot_module().create_backup(tg(request))
    return web.FileResponse(path, headers={"Content-Disposition": f'attachment; filename="{path.name}"'})


async def root(request: web.Request) -> web.Response:
    raise web.HTTPFound("/admin")


def setup_admin(web_app: web.Application, tg_app) -> None:
    web_app["tg"] = tg_app
    web_app["sessions"] = set()
    web_app.middlewares.append(require_login)
    r = web_app.router
    r.add_get("/", root)
    r.add_get("/admin", overview)
    r.add_get("/admin/login", login_form)
    r.add_post("/admin/login", login)
    r.add_post("/admin/logout", logout)
    r.add_get("/admin/products", products)
    r.add_post("/admin/products", create_product)
    r.add_get("/admin/products/{id:\\d+}", product_detail)
    r.add_post("/admin/products/{id:\\d+}", update_product)
    r.add_post("/admin/products/{id:\\d+}/toggle", toggle_product)
    r.add_post("/admin/products/{id:\\d+}/notify", notify_product)
    r.add_post("/admin/products/{id:\\d+}/stock", import_stock)
    r.add_post("/admin/stock/{id:\\d+}/delete", delete_stock)
    r.add_get("/admin/orders", orders)
    r.add_get("/admin/orders.xlsx", export_orders)
    r.add_get("/admin/orders/{id:\\d+}", order_detail)
    r.add_post("/admin/orders/{id:\\d+}/confirm", confirm_order)
    r.add_post("/admin/orders/{id:\\d+}/cancel", cancel_order)
    r.add_post("/admin/orders/{id:\\d+}/resend", resend_order)
    r.add_get("/admin/posts", posts)
    r.add_post("/admin/posts", create_post)
    r.add_post("/admin/posts/test", test_post)
    r.add_post("/admin/posts/{id:\\d+}/cancel", cancel_post)
    r.add_get("/admin/uploads/{name}", upload_file)
    r.add_post("/admin/settings/auto-restock", set_auto_restock)
    r.add_get("/admin/backup.zip", download_backup)
